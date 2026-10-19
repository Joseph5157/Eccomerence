from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Medistore Testing Deployment Plan.docx"

NAVY = "17324D"
PALE_BLUE = "EEF4F8"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(92, 102, 112)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def keep_paragraph(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, (label, width) in enumerate(zip(headers, widths)):
        cell = header_row.cells[index]
        set_width(cell, width)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.3)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[index]
            set_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            run = p.add_run(str(value))
            run.font.size = Pt(9.1)
            run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    keep_paragraph(p, keep_next=True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    keep_paragraph(p)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        keep_paragraph(p)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        keep_paragraph(p)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_width(cell, 6.65)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=140, start=160, bottom=140, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.02
    for index, line in enumerate(lines):
        if index:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8.8)
        run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True

    for name, size, before, after in (
        ("Heading 1", 16, 16, 7),
        ("Heading 2", 12.5, 11, 5),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = TEXT
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    run = p.add_run("Medistore Testing Deployment Plan  |  October 2026")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    add_footer(section)
    configure_styles(doc)

    title = doc.add_paragraph("Medistore Testing Deployment Plan", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(15)
    metadata_run = metadata.add_run("Hostinger Cloud application hosting with Railway data services")
    metadata_run.font.size = Pt(11.5)
    metadata_run.font.color.rgb = MUTED
    metadata_run.italic = True

    add_body(
        doc,
        "For testing, deploy the Next.js storefront and Medusa backend as separate Node.js applications on Hostinger Cloud. Provision PostgreSQL and Redis in one Railway project, then connect the Hostinger backend through Railway public TCP endpoints. Cloudflare is optional for DNS and proxying; it is not the database provider in this plan.",
    )
    add_body(
        doc,
        "This plan gives the setup order, production environment variables, repository changes, and checks required to move the current local application to the testing environment.",
    )

    add_heading(doc, "Architecture", 1)
    add_table(
        doc,
        ["Service", "Platform", "Public address", "Purpose"],
        [
            ["Storefront", "Hostinger Cloud", "store.yourdomain.com", "Next.js customer website"],
            ["Backend and admin", "Hostinger Cloud", "api.yourdomain.com", "Medusa API and admin at /app"],
            ["PostgreSQL", "Railway", "Public TCP endpoint", "Products, users, orders, and application data"],
            ["Redis", "Railway", "Public TCP endpoint", "Events, workflows, locks, and caching"],
            ["DNS and CDN", "Cloudflare optional", "Your domain zones", "DNS, TLS proxying, and edge caching"],
        ],
        [1.25, 1.25, 1.85, 2.30],
    )
    add_body(
        doc,
        "Only the backend connects to PostgreSQL and Redis. The browser calls the Medusa API over HTTPS and must never receive database credentials.",
        bold_lead="Only the backend connects to PostgreSQL and Redis. ",
    )

    add_heading(doc, "Account and domain requirements", 1)
    add_bullets(
        doc,
        [
            "One Hostinger Cloud plan that supports two Node.js applications and custom domains.",
            "One Railway project containing PostgreSQL and Redis services.",
            "A domain with DNS control. Cloudflare may manage DNS, but Hostinger DNS also works.",
            "Fresh production secrets and a new admin account. Do not reuse local testing credentials.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Railway setup", 1)
    add_numbered(
        doc,
        [
            "Create one Railway project for the testing environment.",
            "Add a PostgreSQL service and a Redis service to that project.",
            "For each service, open Settings, then Networking, and enable Public Access.",
            "Copy the PostgreSQL DATABASE_PUBLIC_URL and the Redis REDIS_PUBLIC_URL from Railway.",
            "Keep both public URLs private. Store them only in the Hostinger backend environment variables.",
            "Place Railway and Hostinger services in nearby regions when possible to reduce latency.",
        ],
    )
    add_body(
        doc,
        "Do not use Railway internal DATABASE_URL or REDIS_URL values from Hostinger. Those internal addresses resolve only inside Railway private networking. Public TCP traffic also creates Railway network egress charges.",
        bold_lead="Do not use Railway internal DATABASE_URL or REDIS_URL values from Hostinger. ",
    )

    add_heading(doc, "Hostinger backend deployment", 1)
    add_body(
        doc,
        "Create a Node.js application for the backend repository path. Configure the application to build and start Medusa in production mode, bind to Hostinger's assigned port, and serve the admin interface at https://api.yourdomain.com/app.",
    )
    add_heading(doc, "Backend environment variables", 2)
    add_code_block(
        doc,
        [
            "DATABASE_URL=<Railway DATABASE_PUBLIC_URL>",
            "REDIS_URL=<Railway REDIS_PUBLIC_URL>",
            "STORE_CORS=https://store.yourdomain.com",
            "ADMIN_CORS=https://api.yourdomain.com",
            "AUTH_CORS=https://store.yourdomain.com,https://api.yourdomain.com",
            "JWT_SECRET=<new-random-secret>",
            "COOKIE_SECRET=<new-random-secret>",
            "NODE_ENV=production",
        ],
    )
    add_body(
        doc,
        "Generate long random values for JWT_SECRET and COOKIE_SECRET. Never commit these values or the Railway public URLs to Git.",
    )

    add_heading(doc, "Database initialization", 2)
    add_numbered(
        doc,
        [
            "Run Medusa database migrations using the Hostinger backend environment.",
            "Run the project seed process only if the testing environment needs demo products and categories.",
            "Create a new production admin user after migrations complete.",
            "Create a production publishable API key and associate it with the required sales channel.",
            "Restart the backend and confirm the health endpoint and /app load over HTTPS.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Hostinger storefront deployment", 1)
    add_body(
        doc,
        "Create a second Node.js application for the Next.js storefront. Its backend URL must use the public HTTPS Medusa domain, not localhost and not a Railway database address.",
    )
    add_heading(doc, "Storefront environment variables", 2)
    add_code_block(
        doc,
        [
            "NEXT_PUBLIC_MEDUSA_BACKEND_URL=https://api.yourdomain.com",
            "NEXT_PUBLIC_MEDUSA_PUBLISHABLE_KEY=<production-publishable-key>",
        ],
    )
    add_body(
        doc,
        "Redeploy the storefront whenever a NEXT_PUBLIC variable changes because Next.js includes public variables in the browser build.",
    )

    add_heading(doc, "DNS and HTTPS", 1)
    add_table(
        doc,
        ["Hostname", "Target", "Expected result"],
        [
            ["store.yourdomain.com", "Hostinger storefront application", "Customer website opens over HTTPS"],
            ["api.yourdomain.com", "Hostinger backend application", "Medusa API and /app open over HTTPS"],
        ],
        [1.85, 2.30, 2.50],
    )
    add_body(
        doc,
        "If Cloudflare manages DNS, begin with DNS-only records while validating Hostinger routing and certificates. Enable proxying after both applications work correctly. Do not proxy PostgreSQL or Redis through standard Cloudflare DNS records.",
    )

    add_heading(doc, "Repository preparation", 1)
    add_bullets(
        doc,
        [
            "Make the Medusa server listen on Hostinger's PORT value and on a network-accessible interface.",
            "Configure Medusa Redis modules for event bus, workflow execution, distributed locking, and caching. The current local fallback uses in-memory substitutes and is unsuitable for a deployed multi-process service.",
            "Add explicit build, start, migration, and seed scripts for Hostinger deployment.",
            "Enable secure cookies and confirm cross-domain CORS for the storefront and admin domains.",
            "Keep existing static demo product images in the storefront build. Add object storage later if administrators must upload product images through the admin panel.",
        ],
    )

    add_heading(doc, "Deployment order", 1)
    add_numbered(
        doc,
        [
            "Create Railway PostgreSQL and Redis services.",
            "Enable Railway Public Access and copy the public connection URLs.",
            "Prepare and deploy the Medusa backend on Hostinger.",
            "Add backend environment variables and run migrations.",
            "Create the admin user, publishable API key, and required sales channel association.",
            "Deploy the Next.js storefront on Hostinger with its production variables.",
            "Connect custom domains, confirm HTTPS, and complete the smoke tests.",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "Validation checklist", 1)
    add_table(
        doc,
        ["Area", "Test", "Expected result"],
        [
            ["Storefront", "Open home and product pages", "Pages load without browser console errors"],
            ["Catalog", "Browse Vitamins, Syrups, and Supplements on mobile", "Category cards remain small and horizontal"],
            ["Images", "Open seeded products", "Demo product images load correctly"],
            ["Customer auth", "Register, log in, refresh, and log out", "Session persists securely and logout clears it"],
            ["Admin", "Open /app and sign in", "New production admin can access the dashboard"],
            ["Database", "Create or update a product", "Changes remain after backend restart"],
            ["Redis", "Run a workflow and restart the backend", "No in-memory fallback warnings appear"],
            ["CORS", "Use storefront and admin from their public domains", "API requests succeed without CORS errors"],
            ["Secrets", "Review browser bundle and repository", "No database URLs, Redis URLs, or server secrets are exposed"],
        ],
        [1.15, 2.55, 2.95],
    )

    add_heading(doc, "Testing limitations and production notes", 1)
    add_bullets(
        doc,
        [
            "Public Railway database access is acceptable for this testing setup but increases exposure. Rotate credentials immediately if they are shared or leaked.",
            "Cross-provider database traffic adds latency and egress cost. A future production setup should place application and data services in the same private network when practical.",
            "Use a dedicated object-storage service for admin uploads before relying on user-managed product media in production.",
            "Add automated backups, monitoring, error reporting, and recovery procedures before accepting real customer orders.",
        ],
    )

    add_heading(doc, "Reference documentation", 1)
    references = [
        ("Railway PostgreSQL", "https://docs.railway.com/databases/postgresql"),
        ("Railway Redis", "https://docs.railway.com/databases/redis"),
        ("Hostinger Node.js deployment", "https://www.hostinger.com/support/how-to-deploy-a-nodejs-website-in-hostinger/"),
        ("Hostinger environment variables", "https://www.hostinger.com/support/how-to-add-environment-variables-during-node-js-application-deployment/"),
        ("Medusa deployment", "https://docs.medusajs.com/learn/deployment/general"),
    ]
    for label, url in references:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)
        keep_paragraph(p)

    doc.core_properties.title = "Medistore Testing Deployment Plan"
    doc.core_properties.subject = "Hostinger Cloud and Railway testing deployment"
    doc.core_properties.author = "Medistore"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
